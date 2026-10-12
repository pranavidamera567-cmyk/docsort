"""Text extraction for uploaded files."""
import io

import docx
from pypdf import PdfReader

ALLOWED = {"pdf", "docx", "txt"}


class ExtractionError(Exception):
    pass


def file_type(filename):
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED:
        raise ExtractionError(f"Unsupported file type '.{ext}'. Allowed: {', '.join(sorted(ALLOWED))}")
    return ext


def extract_text(filename, data):
    kind = file_type(filename)
    try:
        if kind == "pdf":
            text = "\n".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(data)).pages)
        elif kind == "docx":
            d = docx.Document(io.BytesIO(data))
            cells = [c.text for t in d.tables for row in t.rows for c in row.cells]
            text = "\n".join([p.text for p in d.paragraphs] + cells)
        else:
            text = data.decode("utf-8", errors="replace")
    except Exception as e:  # corrupt or encrypted files raise assorted library errors
        raise ExtractionError(f"Could not read {kind.upper()} file: {e}") from e
    text = text.strip()
    if not text:
        raise ExtractionError("No extractable text found (scanned or empty document?)")
    return text
